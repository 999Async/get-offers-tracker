"""Opt-in real Docling conversion with synthetic files and offline model weights."""

import os

import pytest

from getoffers_agent.knowledge.contracts import KnowledgeConfig
from getoffers_agent.knowledge.parsing import DocumentParser, make_units

pytestmark = pytest.mark.skipif(
    not os.environ.get("GETOFFERS_DOCUMENT_INTEGRATION"),
    reason="requires Docling and cached PDF layout model",
)


def simple_pdf(text):
    stream = f"BT /F1 14 Tf 50 750 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    raw = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(raw))
        raw += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(raw)
    raw += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(
        f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]
    )
    return raw + f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


def test_real_docx_heading_table_and_chinese(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("项目经历", level=1)
    doc.add_paragraph("实现 Python Agent 检索评测。")
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("指标", "结果"), ("召回率", "已评测")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    path = tmp_path / "resume.docx"
    doc.save(path)
    canonical = DocumentParser(KnowledgeConfig()).parse(path.read_bytes(), "docx")
    assert canonical.nodes[0].kind == "heading"
    assert any("检索评测" in n.text for n in canonical.nodes)
    assert any(n.kind == "table" and n.table_cells for n in canonical.nodes)
    assert all(n.locator.source_ref for n in canonical.nodes)


def test_real_pdf_page_locator_and_citation():
    raw = simple_pdf("Python retrieval evaluation project")
    config = KnowledgeConfig()
    canonical = DocumentParser(config).parse(raw, "pdf")
    assert any("retrieval evaluation" in node.text for node in canonical.nodes)
    assert all(node.locator.page == 1 and node.locator.bbox for node in canonical.nodes)
    for unit in make_units(canonical, "alice", "d", "v", config):
        node = next(n for n in canonical.nodes if n.node_id == unit.locator.node_id)
        assert node.text[unit.locator.start : unit.locator.end] == unit.text
